from datetime import date, timedelta
from decimal import Decimal
import asyncio
import io
import zipfile
import httpx
import pytest
from docx import Document
from sqlalchemy import select
from bp_bot.ai import parse_fields, AI
from bp_bot.config import Settings
from bp_bot.domain import BusinessError
from bp_bot.rates import CBR
from bp_bot.documents import validate_letterhead, build_template
from bp_bot.storage import LocalStorage, valid_key
from bp_bot.jobs import Worker
from bp_bot.models import Job, now
from test_quotes import patch


def test_cbr_nominal_and_date():
    xml = b'<ValCurs Date="28.10.2026"><Valute><CharCode>CNY</CharCode><Nominal>10</Nominal><Value>123,45</Value></Valute></ValCurs>'
    rate, day = CBR(httpx.MockTransport(lambda _: httpx.Response(200, content=xml))).rate(
        "CNY", date(2026, 10, 29)
    )
    assert rate == Decimal("12.345") and day == date(2026, 10, 28)


def test_cbr_failure_does_not_use_fake_rate():
    cbr = CBR(httpx.MockTransport(lambda _: httpx.Response(503)))
    with pytest.raises(BusinessError):
        cbr.rate("EUR", date(2026, 10, 29))


def test_structured_input_never_accepts_prices():
    p = parse_fields("клиент=ООО A&B; позиции=A:2,B:1; поставка=10 дней; до=2026-11-05")
    assert p.items[0].quantity == 2
    assert p.client == "ООО A&B"
    with pytest.raises(BusinessError):
        parse_fields("клиент=Заря; цена=1")


def test_ai_without_key_has_useful_fallback():
    settings = Settings(
        bot_token="fake", database_url="sqlite://", app_env="development", storage_backend="local"
    )
    ai = AI(settings)
    assert asyncio.run(ai.parse("клиент=Заря", date(2026, 10, 29))).client == "Заря"
    with pytest.raises(BusinessError, match="AI пока не настроен"):
        asyncio.run(ai.parse("Сделай КП", date(2026, 10, 29)))


def test_production_requires_durable_storage():
    with pytest.raises(ValueError):
        Settings(bot_token="fake", database_url="sqlite://", app_env="production", storage_backend="local")


def test_letterhead_security_and_preservation(tmp_path):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "ООО АТехника"
    raw = io.BytesIO()
    d.save(raw)
    validate_letterhead(raw.getvalue())
    target = tmp_path / "template.docx"
    build_template(raw.getvalue(), target)
    assert Document(target).sections[0].header.paragraphs[0].text == "ООО АТехника"
    bad = io.BytesIO()
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("word/document.xml", "{{ cycler.__init__ }}")
    with pytest.raises(BusinessError):
        validate_letterhead(bad.getvalue())


def test_private_storage_prevents_traversal(tmp_path):
    s = LocalStorage(tmp_path)
    s.put("q/1.pdf", b"pdf", "application/pdf")
    assert s.get("q/1.pdf") == b"pdf"
    with pytest.raises(ValueError):
        valid_key("../secret")


def test_job_recovery_and_retrieval(service, tmp_path):
    q = service.create(2, patch(), "a")["id"]
    service.finish(2, q, 1)
    service.request_files(2, q, 1)
    store = LocalStorage(tmp_path)
    store.put("letterheads/test.docx", b"head", "application/octet-stream")
    with service.sessions.begin() as s:
        job = s.scalar(select(Job))
        job.state, job.lease_until = "running", now() - timedelta(seconds=1)
    worker = Worker(service.sessions, store, renderer=lambda *_: (b"docx", b"%PDF-1.7"))
    assert worker.once()
    assert not worker.once()
    result = service.request_files(2, q, 1)
    assert result["state"] == "done"
    assert store.get(result["pdf_key"]) == b"%PDF-1.7"


def test_job_failure_is_bounded_and_retryable(service, tmp_path):
    q = service.create(2, patch(), "a")["id"]
    service.finish(2, q, 1)
    service.request_files(2, q, 1)
    worker = Worker(service.sessions, LocalStorage(tmp_path))
    for _ in range(3):
        assert worker.once()
    assert not worker.once()
    with service.sessions() as s:
        job = s.scalar(select(Job))
        assert job.state == "failed" and job.attempts == 3
    assert service.request_files(2, q, 1)["state"] == "queued"


def test_malformed_cbr_xml_fails_closed():
    cbr = CBR(httpx.MockTransport(lambda _: httpx.Response(200, content=b"<broken")))
    with pytest.raises(BusinessError, match="Не удалось получить курс"):
        cbr.rate("EUR", date(2026, 10, 29))


def test_transcription_catalog_context(service, monkeypatch):
    from pathlib import Path
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from bp_bot.models import Product

    with service.sessions.begin() as session:
        for n in range(40):
            session.add(
                Product(
                    sku=f"HINT{n}",
                    name=f"GRITAN БНП {300 + n}",
                    configuration="PRIVATE CONFIG",
                    price=123456,
                    currency="RUB",
                    supply="stock",
                    price_date="2026-01-01",
                    valid_until="2026-12-31",
                    vat_rate=22,
                    active=True,
                )
            )
        session.get(Product, "A").active = False
    prompt = service.transcription_prompt(2)
    assert "БНП 339" in prompt  # Not limited to the 30-row /catalog page.
    assert prompt.count("Трактор 100") == 1
    assert "Бункер 401" not in prompt
    assert "123456" not in prompt and "PRIVATE CONFIG" not in prompt
    assert len(prompt) <= 6000
    with pytest.raises(BusinessError):
        service.transcription_prompt(999)

    settings = Settings(
        bot_token="123456:TEST",
        database_url="sqlite://",
        app_env="development",
        storage_backend="local",
        openai_api_key="",
        openai_transcribe_model="gpt-4o-mini-transcribe",
    )
    ai = AI(settings)
    create = AsyncMock(return_value=SimpleNamespace(text="Гритан БНП 301"))
    ai.client = SimpleNamespace(audio=SimpleNamespace(transcriptions=SimpleNamespace(create=create)))

    def convert(args, **kwargs):
        Path(args[-1]).write_bytes(b"wav")

    monkeypatch.setattr("bp_bot.ai.subprocess.run", convert)
    assert asyncio.run(ai.transcribe(b"ogg", prompt=prompt)) == "Гритан БНП 301"
    assert create.call_args.kwargs["prompt"] == prompt
    assert create.call_args.kwargs["language"] == "ru"
    with service.sessions.begin() as session:
        for p in session.scalars(select(Product)):
            p.active = False
    assert service.transcription_prompt(2) == ""
